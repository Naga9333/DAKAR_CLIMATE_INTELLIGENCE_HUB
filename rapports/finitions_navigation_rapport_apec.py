# -*- coding: utf-8 -*-
"""
Finitions de navigation et de conformité :
1. Tri alphabétique des tables Sigles et Glossaire
2. Signets sur chaque titre + hyperliens dans la TDM (navigation Word et PDF)
"""
import re
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX = "/tmp/doc/RAPPORT_INSTITUTIONNEL_JOURNEES_SCIENTIFIQUES_APEC_2026_PROFESSIONNEL_V10.docx"
d = docx.Document(DOCX)

# ------------------------------------------------ 1. Tri sigles / glossaire
def sort_table_rows(tbl):
    rows = tbl._tbl.findall(qn("w:tr"))
    header, body = rows[0], rows[1:]
    def key(tr):
        first_cell = tr.find(qn("w:tc"))
        return re.sub(r"\s+", " ", "".join(first_cell.itertext())).strip().upper()
    for tr in sorted(body, key=key):
        tbl._tbl.append(tr)  # déplace en fin, dans l'ordre trié

for t in d.tables:
    if len(t.columns) == 2 and t.rows[0].cells[0].text.strip() in ("Sigle", "Notion"):
        sort_table_rows(t)

# ------------------------------------------------ 2. Signets sur les titres
existing_ids = [int(b.get(qn("w:id"))) for b in d.element.body.iter(qn("w:bookmarkStart"))
                if (b.get(qn("w:id")) or "").isdigit()]
next_id = max(existing_ids, default=0) + 1

heading_ps = [p for p in d.paragraphs
              if p.style.name in ("Heading 1", "Heading 2", "Heading 3") and p.text.strip()]

names = []
for k, p in enumerate(heading_ps):
    name = f"_Toc{100001 + k}"
    names.append(name)
    bid = str(next_id + k)
    bs = OxmlElement("w:bookmarkStart")
    bs.set(qn("w:id"), bid); bs.set(qn("w:name"), name)
    be = OxmlElement("w:bookmarkEnd")
    be.set(qn("w:id"), bid)
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(bs)
    else:
        p._p.insert(0, bs)
    p._p.append(be)

# ------------------------------------------------ 3. Hyperliens dans la TDM
toc_ps = [p for p in d.paragraphs if p.style.name in ("TOC 1", "TOC 2", "TOC 3")]
assert len(toc_ps) == len(names), (len(toc_ps), len(names))

FIELD_TAGS = (qn("w:fldChar"), qn("w:instrText"))
for p, name in zip(toc_ps, names):
    entry_runs = []
    for r in p._p.findall(qn("w:r")):
        if any(r.find(tag) is not None for tag in FIELD_TAGS):
            continue  # runs de champ : restent hors hyperlien
        entry_runs.append(r)
    if not entry_runs:
        continue
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("w:anchor"), name)
    hl.set(qn("w:history"), "1")
    entry_runs[0].addprevious(hl)
    for r in entry_runs:
        hl.append(r)

d.save(DOCX)
print("OK : tables triées,", len(names), "signets et hyperliens posés")
