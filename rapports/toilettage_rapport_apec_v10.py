# -*- coding: utf-8 -*-
"""
Toilettage et mise aux normes internationales du rapport institutionnel APEC.
- Format A4, marges 2,5 cm, langue fr-FR
- Typographie française (apostrophes, espaces insécables, ligatures, capitales accentuées, n°, tirets)
- Table des matières automatique
- Styles de titres harmonisés, veuves/orphelines, titres solidaires
- 83 tableaux harmonisés (en-têtes, bordures, encadrés)
- Page de garde, en-tête / pied de page, métadonnées
"""
import re
import copy
import docx
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/root/.claude/uploads/115309f9-d71e-5fd4-908e-e461d14f0ee1/f8ae7c13-RAPPORT_INSTITUTIONNEL_JOURNEES_SCIENTIFIQUES_APEC_2026_PROFESSIONNEL_V9.docx"
DST = "/tmp/doc/RAPPORT_INSTITUTIONNEL_JOURNEES_SCIENTIFIQUES_APEC_2026_PROFESSIONNEL_V10.docx"

NBSP = " "
BLEU_FONCE = "1F4E79"
BLEU_MOYEN = "2E74B5"
GRIS_TEXTE = "595959"
GRIS_BORD = "BFBFBF"
FOND_ENCADRE = "F2F6FA"

d = docx.Document(SRC)

# ---------------------------------------------------------------- 1. Texte
ACCENTS_CAPITALES = [
    (r"\bEtat(s?)\b", r"État\1"),
    (r"\bEtatique", "Étatique"),
    (r"\bElabor", "Élabor"),
    (r"\bEtud", "Étud"),
    (r"\bEvit", "Évit"),
    (r"\bEnergie", "Énergie"),
    (r"\bEvénement", "Événement"),
    (r"\bEvenement", "Événement"),
    (r"\bEchange", "Échange"),
    (r"\bEchelle", "Échelle"),
    (r"\bEconomi", "Économi"),
    (r"\bEquip", "Équip"),
    (r"\bEquité", "Équité"),
    (r"\bEgalité", "Égalité"),
    (r"\bEvalu", "Évalu"),
    (r"\bEvolution", "Évolution"),
    (r"\bEtabli", "Établi"),
    (r"\bEmett", "Émett"),
    (r"\bEmission", "Émission"),
]

def normalize(t):
    if not t:
        return t
    # apostrophe typographique
    t = t.replace("'", "’")
    # ligatures
    for a, b in (("oeuvre", "œuvre"), ("Oeuvre", "Œuvre"), ("coeur", "cœur"),
                 ("Coeur", "Cœur"), ("noeud", "nœud"), ("soeur", "sœur"),
                 ("voeu", "vœu")):
        t = t.replace(a, b)
    # capitales accentuées
    for pat, rep in ACCENTS_CAPITALES:
        t = re.sub(pat, rep, t)
    # À en début de phrase
    t = re.sub(r"(^|[.!?:;]\s+)A\s", r"\1À ", t)
    # numéros de lois et décrets
    t = re.sub(r"\b([Ll]oi|[Dd]écret)\s+no\s+", r"\1 n°" + NBSP, t)
    t = re.sub(r"\bno\s+(\d{4}-\d+)", r"n°" + NBSP + r"\1", t)
    # tiret demi-cadratin
    t = t.replace(" - ", " – ")
    # points de suspension
    t = t.replace("...", "…")
    # guillemets français (paires complètes uniquement)
    t = re.sub(r'"([^"\n]+)"', "«" + NBSP + r"\1" + NBSP + "»", t)
    # espaces multiples
    t = re.sub(r"  +", " ", t)
    # espaces insécables (typographie française)
    t = re.sub(r" ([:;!?»])", NBSP + r"\1", t)
    t = re.sub(r"« ", "«" + NBSP, t)
    t = re.sub(r"(\d) %", r"\1" + NBSP + "%", t)
    t = re.sub(r"(\d) (FCFA|CFA|milliards?|millions?)\b", r"\1" + NBSP + r"\2", t)
    return t

for el in d.element.body.iter(qn("w:t")):
    el.text = normalize(el.text)

# ---------------------------------------------------------------- 2. Langue
def set_lang(rpr_parent):
    rpr = rpr_parent.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_parent.append(rpr)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "fr-FR")
    lang.set(qn("w:eastAsia"), "fr-FR")

styles_el = d.styles.element
for rprd in styles_el.iter(qn("w:rPrDefault")):
    set_lang(rprd)

# ---------------------------------------------------------------- 3. Mise en page A4
sec = d.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, attr, Cm(2.5))
sec.header_distance = Cm(1.25)
sec.footer_distance = Cm(1.25)
sec.different_first_page_header_footer = True  # page de garde sans en-tête/pied

# ---------------------------------------------------------------- 4. Styles
def style_ppr(st):
    return st.element.get_or_add_pPr()

normal = d.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
pf = normal.paragraph_format
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pf.space_after = Pt(8)
pf.space_before = Pt(0)
pf.line_spacing = 1.15
pf.widow_control = True

h1 = d.styles["Heading 1"]
h1.font.name = "Calibri"
h1.font.size = Pt(16)
h1.font.bold = True
h1.font.color.rgb = RGBColor.from_string(BLEU_FONCE)
h1.paragraph_format.space_before = Pt(0)
h1.paragraph_format.space_after = Pt(14)
h1.paragraph_format.keep_with_next = True
h1.paragraph_format.page_break_before = True
# filet sous les titres de niveau 1
ppr = style_ppr(h1)
pbdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "8")
bottom.set(qn("w:space"), "4")
bottom.set(qn("w:color"), BLEU_MOYEN)
pbdr.append(bottom)
ppr.append(pbdr)

h2 = d.styles["Heading 2"]
h2.font.name = "Calibri"
h2.font.size = Pt(13)
h2.font.bold = True
h2.font.color.rgb = RGBColor.from_string(BLEU_MOYEN)
h2.paragraph_format.space_before = Pt(14)
h2.paragraph_format.space_after = Pt(6)
h2.paragraph_format.keep_with_next = True
h2.paragraph_format.page_break_before = False

h3 = d.styles["Heading 3"]
h3.font.name = "Calibri"
h3.font.size = Pt(11.5)
h3.font.bold = True
h3.font.color.rgb = RGBColor.from_string("44546A")
h3.paragraph_format.space_before = Pt(10)
h3.paragraph_format.space_after = Pt(4)
h3.paragraph_format.keep_with_next = True

lb = d.styles["List Bullet"]
lb.font.name = "Calibri"
lb.font.size = Pt(11)
lb.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
lb.paragraph_format.space_after = Pt(4)
lb.paragraph_format.line_spacing = 1.15

# style du titre de la table des matières (hors champ TDM)
toc_h = d.styles.add_style("Titre TDM", WD_STYLE_TYPE.PARAGRAPH)
toc_h.base_style = h1
toc_h.quick_style = True
ol = OxmlElement("w:outlineLvl")
ol.set(qn("w:val"), "9")
style_ppr(toc_h).append(ol)

# ---------------------------------------------------------------- 5. Page de garde
def para_by_text(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None

p_sur = para_by_text("RAPPORT INSTITUTIONNEL")
p_titre = para_by_text("Journées scientifiques de contribution")
p_sous = para_by_text("Financer, gouverner et transformer")

if p_sur is not None:
    p_sur.paragraph_format.space_before = Pt(160)
    p_sur.paragraph_format.space_after = Pt(18)
    p_sur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p_sur.runs:
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(BLEU_MOYEN)
        rpr = r._r.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), "60")  # lettres espacées
        rpr.append(sp)

if p_titre is not None:
    p_titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_titre.paragraph_format.space_after = Pt(14)
    for r in p_titre.runs:
        r.font.size = Pt(26)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(BLEU_FONCE)

if p_sous is not None:
    p_sous.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sous.paragraph_format.space_after = Pt(48)
    for r in p_sous.runs:
        r.font.size = Pt(14)
        r.font.italic = True
        r.font.color.rgb = RGBColor.from_string(GRIS_TEXTE)

# ---------------------------------------------------------------- 6. Table des matières
def make_run(text=None, fld=None, instr=None, italic=False, color=None, size=None):
    r = OxmlElement("w:r")
    if italic or color or size:
        rpr = OxmlElement("w:rPr")
        if italic:
            rpr.append(OxmlElement("w:i"))
        if color:
            c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
        if size:
            sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size * 2)); rpr.append(sz)
        r.append(rpr)
    if fld:
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), fld); r.append(fc)
    if instr is not None:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r.append(it)
    if text is not None:
        t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text; r.append(t)
    return r

toc_heading = None
for p in d.paragraphs:
    if p.text.strip() == "Table des matières" and p.style.name == "Heading 1":
        toc_heading = p
        break

existing_toc = [fs for fs in d.element.body.iter(qn("w:fldSimple"))
                if "TOC" in (fs.get(qn("w:instr")) or "")]
if toc_heading is not None:
    toc_heading.style = toc_h
    if not existing_toc:
        toc_p = OxmlElement("w:p")
        toc_p.append(make_run(fld="begin"))
        toc_p.append(make_run(instr=' TOC \\o "1-3" \\h \\z \\u '))
        toc_p.append(make_run(fld="separate"))
        toc_p.append(make_run(
            text="La table des matières est générée automatiquement : clic droit puis « Mettre à jour les champs ».",
            italic=True, color=GRIS_TEXTE))
        toc_p.append(make_run(fld="end"))
        toc_heading._p.addnext(toc_p)

# mise à jour automatique des champs à l'ouverture + césure
settings = d.settings.element
upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
settings.append(upd)
hyph = OxmlElement("w:autoHyphenation"); hyph.set(qn("w:val"), "true")
settings.append(hyph)

# ---------------------------------------------------------------- 7. En-tête / pied de page
header = sec.header
header.is_linked_to_previous = False
hp = header.paragraphs[0]
hp.text = ""
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = hp.add_run("Journées scientifiques de contribution de l’APEC – Ville de Dakar")
run.font.size = Pt(9)
run.font.color.rgb = RGBColor.from_string(GRIS_TEXTE)
hppr = hp._p.get_or_add_pPr()
pbdr = OxmlElement("w:pBdr")
bt = OxmlElement("w:bottom")
bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "4")
bt.set(qn("w:space"), "4"); bt.set(qn("w:color"), GRIS_BORD)
pbdr.append(bt)
hppr.append(pbdr)

footer = sec.footer
footer.is_linked_to_previous = False
fp = footer.paragraphs[0]
for r in list(fp._p):
    if r.tag != qn("w:pPr"):
        fp._p.remove(r)
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
def small_gray(r):
    rpr = r.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
    c = OxmlElement("w:color"); c.set(qn("w:val"), GRIS_TEXTE); rpr.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "18"); rpr.append(sz)
    return r
def add_field(p, instr, default):
    p.append(small_gray(make_run(fld="begin")))
    p.append(small_gray(make_run(instr=instr)))
    p.append(small_gray(make_run(fld="separate")))
    p.append(small_gray(make_run(text=default)))
    p.append(small_gray(make_run(fld="end")))

fp._p.append(small_gray(make_run(text="Page ")))
add_field(fp._p, " PAGE ", "1")
fp._p.append(small_gray(make_run(text=" sur ")))
add_field(fp._p, " NUMPAGES ", "1")

# page de garde : en-tête / pied vides
sec.first_page_header.is_linked_to_previous = False
sec.first_page_footer.is_linked_to_previous = False

# ---------------------------------------------------------------- 8. Tableaux
def set_borders(tbl, spec):
    tblpr = tbl._tbl.tblPr
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, (val, sz, color) in spec.items():
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), val); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        borders.append(e)
    tblpr.append(borders)

def set_cell_margins(tbl, top=40, bottom=40, left=100, right=100):
    tblpr = tbl._tbl.tblPr
    old = tblpr.find(qn("w:tblCellMar"))
    if old is not None:
        tblpr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)

def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)

def format_cell_text(cell, size=10, bold=None, color=None, align=None):
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if align is not None:
            p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold
            if color is not None:
                r.font.color.rgb = RGBColor.from_string(color)

def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:tblHeader")) is None:
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)

THIN = ("single", 4, GRIS_BORD)
tables = d.tables
cover_table = tables[0]

for i, tbl in enumerate(tables):
    rows, cols = len(tbl.rows), len(tbl.columns)
    if tbl is cover_table:
        # bloc de métadonnées de la page de garde
        set_borders(tbl, {"top": THIN, "bottom": THIN, "left": THIN,
                          "right": THIN, "insideH": THIN, "insideV": THIN})
        set_cell_margins(tbl, top=60, bottom=60, left=140, right=140)
        tblpr = tbl._tbl.tblPr
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblpr.append(jc)
        for row in tbl.rows:
            shade_cell(row.cells[0], "EDF1F6")
            shade_cell(row.cells[1], "FFFFFF")
            format_cell_text(row.cells[0], size=10, bold=True, color=BLEU_FONCE)
            format_cell_text(row.cells[1], size=10)
    elif rows == 1 and cols == 1:
        # encadré (citation, message, point de vigilance)
        set_borders(tbl, {
            "top": ("single", 4, "D9D9D9"),
            "bottom": ("single", 4, "D9D9D9"),
            "right": ("single", 4, "D9D9D9"),
            "left": ("single", 24, BLEU_FONCE),
        })
        set_cell_margins(tbl, top=80, bottom=80, left=160, right=160)
        cell = tbl.rows[0].cells[0]
        shade_cell(cell, FOND_ENCADRE)
        for j, p in enumerate(cell.paragraphs):
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs:
                r.font.size = Pt(10)
                if j == 0:  # intitulé de l'encadré
                    r.font.bold = True
                    r.font.color.rgb = RGBColor.from_string(BLEU_FONCE)
    else:
        # tableau de données : en-tête uniforme, bordures fines
        set_borders(tbl, {"top": THIN, "bottom": THIN, "left": THIN,
                          "right": THIN, "insideH": THIN, "insideV": THIN})
        set_cell_margins(tbl)
        hdr = tbl.rows[0]
        repeat_header(hdr)
        for c in hdr.cells:
            shade_cell(c, BLEU_FONCE)
            format_cell_text(c, size=10, bold=True, color="FFFFFF")
        for row in tbl.rows[1:]:
            for c in row.cells:
                format_cell_text(c, size=10)

# ---------------------------------------------------------------- 9. Paragraphes vides
body = d.element.body
children = list(body)
P, TBL = qn("w:p"), qn("w:tbl")
def is_empty_p(el):
    if el.tag != P:
        return False
    if "".join(el.itertext()).strip():
        return False
    if el.findall(".//" + qn("w:drawing")) or el.findall(".//" + qn("w:fldChar")):
        return False
    return True

for idx, el in enumerate(children):
    if not is_empty_p(el):
        continue
    prev = children[idx - 1] if idx > 0 else None
    nxt = children[idx + 1] if idx + 1 < len(children) else None
    after_tbl = prev is not None and prev.tag == TBL
    before_tbl = nxt is not None and nxt.tag == TBL
    last_block = nxt is None or nxt.tag == qn("w:sectPr")
    if after_tbl and (before_tbl or last_block):
        # séparateur obligatoire entre deux tableaux : on le réduit
        ppr = el.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr"); el.insert(0, ppr)
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:after"), "60"); sp.set(qn("w:line"), "240")
        sp.set(qn("w:lineRule"), "auto")
        ppr.append(sp)
        rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "12"); rpr.append(sz)
        ppr.append(rpr)
    else:
        body.remove(el)

# ---------------------------------------------------------------- 10. En-tête/pied : typographie
for part in (header, footer):
    for p in part.paragraphs:
        for r in p.runs:
            # ne pas toucher aux runs de champ (PAGE, NUMPAGES)
            if r._r.find(qn("w:fldChar")) is not None or r._r.find(qn("w:instrText")) is not None:
                continue
            for t in r._r.findall(qn("w:t")):
                t.text = normalize(t.text)

# ---------------------------------------------------------------- 11. Notes de bas de page
from lxml import etree
for part in d.part.package.iter_parts():
    if str(part.partname) == "/word/footnotes.xml":
        root = etree.fromstring(part.blob)
        for el in root.iter(qn("w:t")):
            el.text = normalize(el.text)
        part._blob = etree.tostring(root, xml_declaration=True,
                                    encoding="UTF-8", standalone=True)

# ---------------------------------------------------------------- 12. Métadonnées
cp = d.core_properties
cp.title = ("Rapport institutionnel – Journées scientifiques de contribution "
            "de l’APEC 2026")
cp.subject = ("Financer, gouverner et transformer la Ville de Dakar à l’ère des "
              "pôles territoriaux et de l’Acte IV de la décentralisation")
cp.author = "APEC – Ville de Dakar"
cp.last_modified_by = "APEC – Ville de Dakar"
cp.language = "fr-FR"
cp.category = "Rapport institutionnel"
cp.keywords = "Dakar; APEC; décentralisation; Acte IV; GAR; finances locales"

d.save(DST)
print("OK ->", DST)
